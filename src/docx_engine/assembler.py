"""Assemble section JSONs into the final .docx file."""

import json
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from rich.console import Console

from . import styles
from .renderer import render_content_item, set_run_font

console = Console()

# Section ordering for assembly
SECTION_FILE_ORDER = [
    "purpose_scope",        # Sections 1 + 2
    "references",           # Section 3
    "definitions",          # Section 4
    "overview",             # Section 5
    "components",           # Section 6.1
    "integrations",         # Section 6.2
    # feature_6_3_1, feature_6_3_2, etc. — sorted dynamically
    "ai_principles",        # Section 6.4
    "security",             # Section 6.5
    "attachments",          # Section 7
]


def _add_doc_identification(doc: Document, sw_number: str, sw_name: str):
    """Add the document identification header."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run("DOCUMENT IDENTIFICATION:")
    set_run_font(run, styles.HEADING_FONT, Pt(10), bold=True)

    p2 = doc.add_paragraph()
    run1 = p2.add_run(f"{sw_number} {sw_name}")
    set_run_font(run1, styles.HEADING_FONT, Pt(14), bold=True)

    p3 = doc.add_paragraph()
    run2 = p3.add_run("Software Design Specification")
    set_run_font(run2, styles.HEADING_FONT, Pt(14))

    # Add some spacing
    doc.add_paragraph()


def _add_toc_placeholder(doc: Document):
    """Add a Table of Contents placeholder.
    Note: TOC auto-generation requires updating fields in Word.
    We add the field code; user must right-click > Update Field in Word."""
    p = doc.add_paragraph()
    run = p.add_run("TABLE OF CONTENTS")
    set_run_font(run, styles.HEADING_FONT, Pt(12), bold=True)

    # Add TOC field code
    p2 = doc.add_paragraph()
    run = p2.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fldChar1)

    run2 = p2.add_run()
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._element.append(instrText)

    run3 = p2.add_run()
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._element.append(fldChar2)

    run4 = p2.add_run("[Right-click and select 'Update Field' to generate Table of Contents]")
    set_run_font(run4, styles.BODY_FONT, Pt(10))

    run5 = p2.add_run()
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._element.append(fldChar3)

    doc.add_page_break()


def _render_section(doc: Document, section_data, is_top_level: bool = True):
    """Render a single section (or list of sections) into the document."""
    if isinstance(section_data, list):
        for item in section_data:
            _render_section(doc, item, is_top_level)
        return

    if not isinstance(section_data, dict):
        return

    # If it has raw_text (fallback from JSON parse failure), render as paragraph
    if "raw_text" in section_data:
        p = doc.add_paragraph()
        run = p.add_run(section_data["raw_text"][:5000])
        set_run_font(run, styles.BODY_FONT, styles.BODY_SIZE)
        return

    section_number = section_data.get("section_number", "")
    heading_text = section_data.get("heading", "")
    content = section_data.get("content", [])

    # Determine heading level from section number
    if section_number:
        depth = section_number.count(".")
        if depth == 0:
            heading_level = 1  # "1", "2", "3"...
        elif depth == 1:
            heading_level = 2  # "6.1", "6.2"...
        else:
            heading_level = 3  # "6.3.1", "6.3.2"...

        # Add section heading
        full_heading = f"{section_number}  {heading_text}" if heading_text else section_number
        h = doc.add_heading(full_heading, level=heading_level)
        for run in h.runs:
            size_map = {1: styles.HEADING1_SIZE, 2: styles.HEADING2_SIZE, 3: styles.HEADING3_SIZE}
            set_run_font(run, styles.HEADING_FONT, size_map.get(heading_level, styles.HEADING3_SIZE), bold=True)

    # Render content items
    if isinstance(content, list):
        for item in content:
            render_content_item(doc, item, section_number)
    elif isinstance(content, str):
        # Plain text fallback
        p = doc.add_paragraph()
        run = p.add_run(content)
        set_run_font(run, styles.BODY_FONT, styles.BODY_SIZE)


def assemble_docx(
    sections_dir: Path,
    output_path: Path,
    sw_number: str,
    sw_name: str,
    template_path: Optional[Path] = None,
):
    """
    Read all section JSONs and assemble the final .docx.

    Args:
        sections_dir: Directory containing section JSON files
        output_path: Where to save the .docx
        sw_number: Software identifier (e.g., "SW14552")
        sw_name: Software name (e.g., "Advanced Insights Services")
        template_path: Optional custom .docx template
    """
    from typing import Optional

    console.print("\n[bold blue]═══ Assembling .docx ═══[/bold blue]\n")

    # Create document (from template or blank)
    if template_path and template_path.exists():
        doc = Document(str(template_path))
    else:
        doc = Document()
        _setup_default_styles(doc)

    # Document identification
    _add_doc_identification(doc, sw_number, sw_name)

    # Table of Contents
    _add_toc_placeholder(doc)

    # Add "DESIGN" heading before subsections if we have design content
    design_heading_added = False

    # Load and render sections in order
    for section_key in SECTION_FILE_ORDER:
        section_file = sections_dir / f"{section_key}.json"
        if not section_file.exists():
            continue

        # Add "DESIGN" section heading before components
        if section_key == "components" and not design_heading_added:
            h = doc.add_heading("6  DESIGN", level=1)
            for run in h.runs:
                set_run_font(run, styles.HEADING_FONT, styles.HEADING1_SIZE, bold=True)
            design_heading_added = True

        try:
            data = json.loads(section_file.read_text())
            _render_section(doc, data)
        except (json.JSONDecodeError, Exception) as e:
            console.print(f"  [yellow]Warning: Could not render {section_key}: {e}[/yellow]")

        # After integrations, render feature sections
        if section_key == "integrations":
            # Add "Key Features and Functions" heading
            h = doc.add_heading("6.3  Key Features and Functions", level=2)
            for run in h.runs:
                set_run_font(run, styles.HEADING_FONT, styles.HEADING2_SIZE, bold=True)

            # Find and render all feature files in order
            feature_files = sorted(sections_dir.glob("feature_*.json"))
            for feat_file in feature_files:
                try:
                    feat_data = json.loads(feat_file.read_text())
                    _render_section(doc, feat_data)
                except (json.JSONDecodeError, Exception) as e:
                    console.print(f"  [yellow]Warning: Could not render {feat_file.name}: {e}[/yellow]")

    # Save
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    console.print(f"\n[bold green]✓ SDS saved to: {output_path}[/bold green]")

    return output_path


def _setup_default_styles(doc: Document):
    """Set up default styles for a blank document (when no template is used)."""
    style = doc.styles

    # Set default font
    default_font = style["Normal"].font
    default_font.name = styles.BODY_FONT
    default_font.size = styles.BODY_SIZE

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(styles.MARGIN_INCHES)
        section.bottom_margin = Inches(styles.MARGIN_INCHES)
        section.left_margin = Inches(styles.MARGIN_INCHES)
        section.right_margin = Inches(styles.MARGIN_INCHES)
